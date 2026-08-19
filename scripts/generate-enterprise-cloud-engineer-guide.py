#!/usr/bin/env python3
"""Generate the Enterprise Cloud Engineer Azure concepts and Q&A guide."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Enterprise_Cloud_Engineer_Azure_Concepts_and_QA.docx"

BLUE = "0078D4"
DARK_BLUE = "17365D"
LIGHT_BLUE = "DDEBF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction_text, separate, end])


def add_bullets(document, items, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        document.add_paragraph(item, style=style)


def add_numbered(document, items):
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_qa(document, question, answer, number):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Question"]
    paragraph.add_run(f"Q{number}. {question}").bold = True
    answer_paragraph = document.add_paragraph()
    answer_paragraph.style = document.styles["Answer"]
    answer_paragraph.add_run("Answer: ").bold = True
    answer_paragraph.add_run(answer)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 30, DARK_BLUE),
        ("Subtitle", 15, BLUE),
        ("Heading 1", 20, DARK_BLUE),
        ("Heading 2", 15, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    if "Question" not in styles:
        question = styles.add_style("Question", WD_STYLE_TYPE.PARAGRAPH)
    else:
        question = styles["Question"]
    question.font.name = "Aptos"
    question.font.size = Pt(10.5)
    question.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    question.paragraph_format.space_before = Pt(8)
    question.paragraph_format.space_after = Pt(2)
    question.paragraph_format.keep_with_next = True

    if "Answer" not in styles:
        answer = styles.add_style("Answer", WD_STYLE_TYPE.PARAGRAPH)
    else:
        answer = styles["Answer"]
    answer.font.name = "Aptos"
    answer.font.size = Pt(10)
    answer.paragraph_format.left_indent = Inches(0.2)
    answer.paragraph_format.space_after = Pt(7)

    header = section.header.paragraphs[0]
    header.text = "Enterprise Cloud Engineer — Azure Concepts and Design Q&A"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    add_field(footer, "PAGE")
    footer.add_run(" of ")
    add_field(footer, "NUMPAGES")


CONCEPTS = [
    {
        "title": "Azure Landing Zones",
        "definition": (
            "An Azure landing zone is a governed, scalable Azure environment in which "
            "workloads can be deployed safely. It is not a single template or a virtual "
            "network. It is the combination of tenant structure, subscriptions, identity, "
            "networking, policy, security, management, automation, and an operating model. "
            "A platform landing zone supplies shared capabilities; an application landing "
            "zone is the workload subscription or subscription set consumed by a product team."
        ),
        "principles": [
            "Design for scale: use management groups and subscription boundaries instead of managing every resource individually.",
            "Keep the tenant root minimally configured and place an organization-specific intermediate root below it.",
            "Separate platform responsibilities—management, connectivity, identity, and security—from workload subscriptions.",
            "Use policy inheritance for consistent controls and allow controlled exceptions with owners and expiry dates.",
            "Treat subscription onboarding as a product with a request, validation, deployment, outputs, support, and lifecycle.",
            "Use Infrastructure as Code and versioned configuration so every platform change is reviewable and repeatable.",
        ],
        "building_blocks": [
            ("Resource organization", "Management groups, subscriptions, resource groups, naming, tagging, ownership, and lifecycle."),
            ("Identity and access", "Microsoft Entra ID, least-privilege RBAC, privileged access, emergency access, and workload identities."),
            ("Network topology", "Connectivity model, IP addressing, DNS, ingress, egress, private access, segmentation, and hybrid connectivity."),
            ("Security", "Defender for Cloud, secure configuration, data protection, threat detection, and incident response."),
            ("Governance", "Policy definitions, initiatives, assignments, exemptions, cost controls, and compliance reporting."),
            ("Management", "Logs, metrics, alerts, update management, backup, inventory, and operational ownership."),
            ("Platform automation", "Terraform or Bicep, remote state, OIDC, tested modules, promotion, and drift management."),
        ],
        "implementation": [
            "Discover business units, regulatory needs, identity boundaries, network dependencies, current subscriptions, and operating teams.",
            "Choose a management group hierarchy based on control differences, not the organization chart alone.",
            "Define platform subscriptions and ownership for connectivity, management, identity, and security.",
            "Deploy the baseline in audit mode, measure impact, remediate, and then move selected controls to enforcement.",
            "Build a subscription-vending process that records owner, cost, classification, destination, network, budget, and expiration.",
            "Pilot with a sandbox subscription, validate operations, then onboard nonproduction and production workloads in waves.",
        ],
        "pitfalls": [
            "Creating too many management groups, which increases assignment complexity and slows change.",
            "Placing deny policies globally before brownfield impact is known.",
            "Using one subscription for platform and workloads, causing ownership, quota, cost, and blast-radius problems.",
            "Calling a network-only deployment a landing zone while ignoring identity, operations, governance, and lifecycle.",
        ],
        "engineer_qa": [
            (
                "What is the difference between a platform landing zone and an application landing zone?",
                "The platform landing zone hosts shared capabilities such as hub networking, identity services, monitoring, policy, and security operations. An application landing zone is the governed subscription environment consumed by a workload team. The platform team owns the first; the product team usually owns resources in the second under inherited controls.",
            ),
            (
                "How would you onboard an existing subscription into an ALZ?",
                "Inventory resources, owners, RBAC, policy, network dependencies, quotas, and costs first. Run policy impact in audit mode, prepare DNS and routing, establish state ownership, move the subscription to a canary management group, remediate findings, and only then move it to the final Corp or Online group. Keep a rollback and communication plan because the move changes inherited policy and access.",
            ),
            (
                "What makes subscription vending enterprise-ready?",
                "A structured request schema, validation, unique state, idempotent modules, least-privilege automation, reviewed plans, approval gates, tags, budgets, diagnostics, access, network integration, clear outputs, and decommission support. It must also report duration, failures, compliance, and exceptions so the platform can be operated as a service.",
            ),
        ],
        "design_qa": [
            (
                "Design an ALZ for a global regulated company with autonomous business units.",
                "Use a minimal tenant root and an enterprise intermediate root. Put shared Platform, Sandbox, Decommissioned, and Landing Zones branches beneath it; create business-unit children only where policy, sovereignty, identity, or network controls truly differ. Separate regional connectivity subscriptions, centralize security and management evidence, use policy initiatives with region-specific parameters, and vend workload subscriptions with owner, classification, residency, budget, and network profiles.",
            ),
            (
                "When would you create another management group?",
                "Create one when a stable control boundary requires different policy, access delegation, sovereignty, or operating ownership across multiple subscriptions. Do not create one merely for reporting, one application, or a temporary project; tags, Resource Graph, resource groups, and subscriptions handle those needs with less hierarchy complexity.",
            ),
            (
                "How do you introduce ALZ into a brownfield tenant?",
                "Start with inventory and dependency mapping. Deploy definitions without assignments, then non-enforcing assignments at a canary scope, analyze exemptions and remediation, establish the target hierarchy, and move subscriptions in low-risk waves. Preserve emergency access, test DNS and routes, compare compliance before and after, and never combine policy enforcement, network migration, and ownership transfer in one uncontrolled change.",
            ),
        ],
    },
    {
        "title": "Azure Governance",
        "definition": (
            "Azure governance is the system of decision rights, technical controls, evidence, "
            "and lifecycle processes that keeps cloud use aligned with business, security, "
            "compliance, and financial goals. Governance is not only Azure Policy. It combines "
            "resource hierarchy, Policy, RBAC and PIM, tags, locks, budgets, naming, compliance "
            "reporting, exception management, and accountable owners."
        ),
        "principles": [
            "Apply controls at the highest safe scope and inherit them downward.",
            "Separate authorization (RBAC) from configuration compliance (Policy).",
            "Prefer built-in policies and roles; create custom assets only for a documented gap.",
            "Start new controls in audit, measure impact, remediate, canary enforce, and then expand.",
            "Make exemptions explicit, narrow, approved, compensated, and time-bound.",
            "Measure outcomes such as compliance, exception age, permanent privilege, budget variance, and remediation success.",
        ],
        "building_blocks": [
            ("Management groups", "Provide inherited governance and access scopes above subscriptions."),
            ("Azure Policy", "Evaluates resource properties and can audit, deny, modify, append, or deploy supporting configuration."),
            ("Initiatives", "Group related policies into a baseline with shared parameters and control mappings."),
            ("RBAC and PIM", "Control who can perform actions and make privileged access eligible, approved, and time-limited."),
            ("Tags and naming", "Provide ownership, cost, environment, classification, and lifecycle metadata."),
            ("Locks", "Protect selected resources from accidental deletion or modification; they do not replace RBAC or backup."),
            ("Cost Management", "Budgets, alerts, exports, reservations, savings plans, and showback support financial accountability."),
            ("Compliance evidence", "Policy Insights, Resource Graph, logs, change history, exemptions, and release records demonstrate control operation."),
        ],
        "implementation": [
            "Define a control catalog with risk, intent, technical mechanism, owner, evidence, and exception path.",
            "Map each control to a management group or subscription scope and parameterize values such as regions and workspace IDs.",
            "Package related policies into initiatives and include noncompliance messages that tell teams how to remediate.",
            "Use managed identities only for effects that need deployment or modification and grant the exact remediation roles.",
            "Automate policy validation, plan, deployment, compliance reporting, and exemption expiry checks.",
            "Review broad permanent RBAC assignments regularly and use PIM for human privilege.",
        ],
        "pitfalls": [
            "Using Owner everywhere because application role requirements were not designed.",
            "Deploying DeployIfNotExists without the managed identity roles needed for remediation.",
            "Treating a policy-compliant resource as automatically secure without considering runtime threats and data.",
            "Allowing permanent exemptions with no owner, compensating control, or expiry.",
        ],
        "engineer_qa": [
            (
                "What is the difference between Azure Policy and RBAC?",
                "RBAC determines who or what can perform an Azure Resource Manager action at a scope. Policy evaluates whether the resulting resource configuration is allowed or compliant. RBAC might allow a developer to create storage accounts while Policy denies public network access or unapproved regions.",
            ),
            (
                "Explain Audit, Deny, Modify, and DeployIfNotExists.",
                "Audit reports a noncompliant resource without blocking it. Deny blocks a create or update that violates the rule. Modify changes supported request properties and can remediate existing resources. DeployIfNotExists creates related configuration, such as a diagnostic setting, after evaluating the resource and requires a managed identity with appropriate roles.",
            ),
            (
                "How do you safely manage policy exemptions?",
                "Capture the requesting owner, business reason, affected control and scope, risk, compensating control, approver, and expiration. Keep scope as narrow as possible, report approaching expirations, and require reapproval rather than automatic renewal. Exemption data is operational evidence, not an invisible bypass.",
            ),
        ],
        "design_qa": [
            (
                "How would you design a policy rollout for thousands of existing resources?",
                "Inventory current compliance first and group findings by control, owner, and remediation type. Deploy in audit, exclude only justified platform scopes, remediate automatically where safe, enforce in a representative canary, and expand by management group after error, support, and exception metrics are acceptable. Keep an emergency rollback procedure for assignment enforcement.",
            ),
            (
                "Where should governance responsibility sit?",
                "A central governance or security team should define risk and mandatory controls; the platform team should implement reusable policy and evidence automation; workload teams should remediate and own exceptions at their scope. A clear RACI prevents the central team from becoming a ticket bottleneck or product teams from self-approving risk.",
            ),
            (
                "How would you enforce tags without blocking resource deployment unnecessarily?",
                "Inherit stable values such as cost center or owner from the subscription or resource group using Modify where supported. Audit resource-specific values first, provide automation defaults, and deny only fields that are essential at create time and cannot be repaired later. Exclude resource types that do not support tags and test child-resource behavior.",
            ),
        ],
    },
    {
        "title": "Azure Networking",
        "definition": (
            "Azure networking connects users, applications, Azure services, the internet, and "
            "on-premises environments while controlling reachability and traffic inspection. "
            "An enterprise design must address IP allocation, topology, routing, DNS, ingress, "
            "egress, segmentation, private service access, resiliency, observability, and "
            "ownership together. A VNet alone is only one component."
        ),
        "principles": [
            "Allocate nonoverlapping address space from an enterprise IP plan and reserve growth space.",
            "Make packet paths and DNS resolution paths explicit in diagrams and tests.",
            "Default to least reachability; segment by trust and workload need rather than by subnet count alone.",
            "Centralize shared controls where useful without creating a single unscaled bottleneck.",
            "Design ingress, east-west, egress, and management traffic separately.",
            "Use zones and regions deliberately and understand which network resources are regional or global.",
        ],
        "building_blocks": [
            ("Virtual networks and subnets", "Provide private address spaces and workload segmentation."),
            ("NSGs and ASGs", "Filter layer 3/4 traffic and group virtual machine interfaces by application role."),
            ("UDRs and BGP", "Control forwarding and exchange routes with appliances, gateways, and external networks."),
            ("Load balancing", "Azure Load Balancer, Application Gateway, Front Door, and Traffic Manager address different layers and scopes."),
            ("Azure Firewall", "Provides managed stateful filtering, threat intelligence, DNAT/SNAT, application rules, and logging."),
            ("DNS", "Azure DNS, Private DNS, Private Resolver, and custom DNS determine service discovery."),
            ("DDoS protection", "Protects public IP resources at network or platform level and complements application-layer protection."),
            ("Network Watcher", "Flow logs, Connection Monitor, packet capture, topology, and troubleshooting provide operational visibility."),
        ],
        "implementation": [
            "Document source, destination, protocol, port, DNS name, route, inspection point, and owner for important flows.",
            "Create an IP address management plan before building VNets or hybrid links.",
            "Select hub-spoke, Virtual WAN, or another topology based on scale, transit, operating model, and feature needs.",
            "Design DNS resolver locations and forwarding rules before private endpoints are deployed.",
            "Control egress with explicit routes, firewall policy, FQDN rules where suitable, and reliable logging.",
            "Test effective routes, effective security rules, name resolution, failover, and throughput from representative workloads.",
        ],
        "pitfalls": [
            "Overlapping CIDRs that prevent peering or complicate hybrid routing.",
            "Using NSGs as the only security control while uncontrolled egress remains available.",
            "Forcing all traffic through an appliance without validating asymmetric routing, SNAT, scale, or health probes.",
            "Deploying private endpoints before a cross-environment DNS model exists.",
        ],
        "engineer_qa": [
            (
                "What happens when Azure chooses a route?",
                "Azure uses longest-prefix match first. If prefixes are equal, route source priority is generally user-defined route, BGP route, then system route. Engineers must inspect effective routes because peering, gateways, service endpoints, private endpoints, and propagated BGP routes can alter the expected path.",
            ),
            (
                "When should you use Application Gateway versus Front Door?",
                "Application Gateway is a regional layer-7 load balancer inside a VNet and can provide WAF and private frontend access. Front Door is a global edge service for internet-facing HTTP/S acceleration, global routing, CDN features, and WAF. A global application may use Front Door at the edge and Application Gateway or an internal service behind it in each region.",
            ),
            (
                "How do NSGs and Azure Firewall differ?",
                "NSGs are distributed stateful layer 3/4 filters applied to subnets or interfaces. Azure Firewall is a centralized managed firewall with network, application, threat-intelligence, DNAT, SNAT, and richer logging features. They are complementary: NSGs constrain local reachability while the firewall controls inspected transit and egress paths.",
            ),
        ],
        "design_qa": [
            (
                "Design network connectivity for 200 subscriptions across four regions.",
                "Use a governed IP plan and regional connectivity hubs or secured Virtual WAN hubs. Place shared DNS, firewall, gateways, and monitoring in connectivity subscriptions; connect spokes through automated profiles; summarize routes where possible; separate production and nonproduction controls when risk requires it; and design hub and gateway scale units from measured throughput and connection counts.",
            ),
            (
                "How would you design controlled internet egress?",
                "Route workload default traffic to zonal Azure Firewall instances or an approved NVA, use policy collections with application and network rules, provide explicit DNS, log decisions centrally, and remove unintended direct public paths. Account for SNAT port scale, service tags, FQDN limitations, platform dependencies, health checks, and a documented emergency access process.",
            ),
            (
                "How do you troubleshoot a connection failure?",
                "Confirm DNS first, then source and destination addresses, effective NSGs, effective routes, appliance policy, return route, and service-level access settings. Use Connection Troubleshoot, Network Watcher, firewall logs, flow logs, packet capture where supported, and application logs. Diagnose both directions because stateful devices still require a valid return path.",
            ),
        ],
    },
    {
        "title": "Azure Hybrid Cloud",
        "definition": (
            "Hybrid cloud connects and consistently operates resources across Azure, on-premises "
            "datacenters, edge sites, and sometimes other clouds. It includes more than a VPN. "
            "A complete design covers connectivity, routing, DNS, identity, certificate and "
            "secret trust, security, monitoring, patching, inventory, data movement, resilience, "
            "and a migration or steady-state operating model."
        ),
        "principles": [
            "Design hybrid dependencies as production services with owners, capacity, monitoring, and recovery objectives.",
            "Use redundant physical and logical paths and test failover rather than assuming provider redundancy.",
            "Keep identity and DNS available during connectivity failures where business requirements demand it.",
            "Extend governance with Azure Arc where it provides clear management value; do not imply Arc turns a server into an Azure VM.",
            "Choose VPN, ExpressRoute, or both based on bandwidth, latency consistency, availability, security, cost, and recovery needs.",
            "Plan address translation explicitly when overlapping networks cannot be renumbered immediately.",
        ],
        "building_blocks": [
            ("Site-to-site VPN", "Encrypted IPsec connectivity over the internet; appropriate for labs, branches, backup, and many moderate workloads."),
            ("ExpressRoute", "Private provider connectivity with predictable routing and enterprise bandwidth options; encryption must be designed separately where required."),
            ("Virtual WAN", "Managed global transit, branch connectivity, routing, security integration, and hub scale."),
            ("Azure Arc", "Projects servers, Kubernetes, and selected data services into Azure management and governance planes."),
            ("Hybrid DNS", "Private Resolver, conditional forwarders, private zones, and on-premises resolvers support bidirectional names."),
            ("Identity", "Microsoft Entra ID, Entra Connect or Cloud Sync, federation choices, managed identities, and local continuity."),
            ("Operations", "Azure Monitor Agent, Update Manager, Defender for Cloud, automation, inventory, and service health."),
        ],
        "implementation": [
            "Collect application dependency, latency, bandwidth, protocol, IP, DNS, and recovery requirements.",
            "Build redundant gateways and circuits or tunnels in separate failure domains.",
            "Define BGP advertisements, route filters, propagation, summarization, and preferred paths.",
            "Create bidirectional DNS forwarding and test private endpoint names from both Azure and on-premises clients.",
            "Onboard representative Arc resources and verify identity, policy, monitoring, patching, and disconnection behavior.",
            "Run failover exercises for circuit, gateway, DNS, identity, and monitoring paths.",
        ],
        "pitfalls": [
            "Treating ExpressRoute as encrypted by default or as a replacement for all internet connectivity.",
            "Advertising a default route without providing Azure services and control-plane dependencies a valid path.",
            "Ignoring DNS forwarding, which makes reachable private IPs unusable by applications.",
            "Assuming Arc-enabled resources inherit every Azure capability or availability guarantee.",
        ],
        "engineer_qa": [
            (
                "When would you select VPN instead of ExpressRoute?",
                "Use VPN when encrypted internet transport, lower entry cost, rapid setup, or branch connectivity meets requirements. Use ExpressRoute when workloads need higher bandwidth, more predictable latency, private provider routing, or enterprise connectivity SLAs. Many critical environments use ExpressRoute as primary and VPN as a diverse backup after route preference and capacity are tested.",
            ),
            (
                "What does Azure Arc provide for servers?",
                "Arc installs a connected machine agent and creates an Azure resource projection for a non-Azure server. That enables inventory, tags, RBAC, policy guest configuration, monitoring, update management, Defender integration, and extensions where supported. Compute, network, and hardware lifecycle remain with the original environment.",
            ),
            (
                "How should hybrid DNS work with Azure private endpoints?",
                "Azure clients typically resolve through Azure Private DNS zones linked to VNets. On-premises resolvers conditionally forward the relevant private-link namespaces to Azure DNS Private Resolver inbound endpoints; Azure can forward corporate namespaces through outbound endpoints and rulesets. Central registration ownership prevents separate teams from creating conflicting records.",
            ),
        ],
        "design_qa": [
            (
                "Design resilient hybrid connectivity for a critical application.",
                "Use diverse provider paths, redundant zone-redundant gateways where supported, BGP, and capacity for the failure state. Define ExpressRoute and VPN route preference, deploy hybrid DNS resolvers across failure domains, monitor BGP and application transactions, and test actual failover. Include identity, certificate, and data replication dependencies in the recovery design.",
            ),
            (
                "How would you handle overlapping address space after an acquisition?",
                "Create an inventory and long-term renumbering plan first. Use NAT on supported VPN gateways, Virtual WAN, or controlled appliances as a temporary translation boundary; isolate routing domains and publish translated DNS records. Document operational complexity because NAT complicates logs, identity-based rules, troubleshooting, and application allowlists.",
            ),
            (
                "What should be measured for hybrid operations?",
                "Tunnel or circuit availability, BGP state, latency, jitter, packet loss, throughput, gateway utilization, route changes, DNS success, Arc connectivity, agent health, and application transactions. Infrastructure green status is insufficient if users cannot resolve names or complete business calls.",
            ),
        ],
    },
    {
        "title": "Azure Security",
        "definition": (
            "Azure security is a layered program for protecting identity, devices, networks, "
            "applications, data, management planes, and operations. It follows shared "
            "responsibility: Microsoft secures the cloud platform, while customers remain "
            "responsible for identities, configuration, data, code, access, monitoring, and "
            "many workload controls. Zero Trust means verify explicitly, use least privilege, "
            "and assume breach."
        ),
        "principles": [
            "Identity is a primary control plane: require strong authentication, conditional access, PIM, and workload identities.",
            "Eliminate permanent secrets where managed identity, workload identity, or federation is available.",
            "Minimize public exposure and protect unavoidable ingress with WAF, DDoS controls, authentication, and monitoring.",
            "Encrypt data in transit and at rest, then control key custody, rotation, recovery, and access.",
            "Centralize security telemetry while preserving workload context and response ownership.",
            "Design prevention, detection, response, and recovery together.",
        ],
        "building_blocks": [
            ("Microsoft Entra ID", "Authentication, Conditional Access, identity protection, workload identities, and access reviews."),
            ("PIM", "Eligible, approved, time-bound privileged access with audit history."),
            ("Defender for Cloud", "Security posture, recommendations, regulatory views, workload protection plans, and alerts."),
            ("Microsoft Sentinel", "Cloud-native SIEM and SOAR for analytics, incidents, hunting, automation, and investigation."),
            ("Key Vault and Managed HSM", "Secrets, keys, certificates, authorization, private access, logging, and recovery."),
            ("Network security", "Firewall, WAF, DDoS, private access, NSGs, segmentation, and controlled egress."),
            ("Data security", "Classification, encryption, database authorization, immutable backup, and exfiltration controls."),
            ("DevSecOps", "Code, dependency, container, IaC, secret, and artifact scanning with signed release processes."),
        ],
        "implementation": [
            "Inventory human, workload, automation, and emergency identities; remove shared and stale credentials.",
            "Apply Conditional Access with tested emergency exclusions and phishing-resistant authentication for privileged roles.",
            "Build secure service baselines for storage, databases, compute, containers, Key Vault, and networking.",
            "Enable Defender plans based on workload risk and send selected logs and alerts to the security operations platform.",
            "Create incident playbooks with owner, severity, evidence, isolation, credential rotation, recovery, and communication steps.",
            "Test backup restoration, key recovery, compromised identity response, and public exposure detection.",
        ],
        "pitfalls": [
            "Storing client secrets in CI when OIDC federation is available.",
            "Sending every log to Sentinel without retention, cost, detection, and ownership design.",
            "Disabling public access without first fixing DNS and trusted management paths.",
            "Assuming encryption alone protects data while broad identities can still read it.",
        ],
        "engineer_qa": [
            (
                "What is the difference between a managed identity and a service principal?",
                "A managed identity is a Microsoft-managed service principal lifecycle attached to an Azure resource; Azure handles credential issuance and rotation. A general service principal represents an application and may use a certificate, secret, or federated credential. Prefer managed identity for Azure-hosted workloads and federation for external automation such as GitHub Actions.",
            ),
            (
                "How do Defender for Cloud and Sentinel differ?",
                "Defender for Cloud focuses on cloud security posture and workload protection, including recommendations and resource-specific alerts. Sentinel is a SIEM/SOAR that correlates data from Azure and other sources into incidents, hunting, and automated response. Defender alerts are commonly one of Sentinel's data sources.",
            ),
            (
                "What is an enterprise Key Vault design?",
                "Separate vaults by application, environment, region, and blast radius where appropriate; use RBAC or access policy consistently, private endpoints, purge protection, soft delete, diagnostics, and controlled network access. Use managed identities, rotate secrets and keys, restrict data-plane roles, and design regional recovery and dependency behavior.",
            ),
        ],
        "design_qa": [
            (
                "Design secure access for GitHub Actions deploying Azure infrastructure.",
                "Create separate plan and apply identities with federated credentials constrained to the repository, branch, and protected environment subjects. Grant read access to plan and narrowly scoped deployment roles to apply, use Entra-authenticated remote state, require reviewed plans and environment approval, pin actions, scan IaC, and retain immutable workflow and Azure activity evidence.",
            ),
            (
                "How would you secure a public multi-region web application?",
                "Use Front Door with WAF and managed certificates at the global edge, private regional origins where supported, strong application identity, DDoS protection for public regional resources, controlled egress, managed identities, Key Vault, private data services, secure headers, patching, Defender, centralized logs, and tested regional failover. Protect admin paths separately with stronger identity and network restrictions.",
            ),
            (
                "How do you design break-glass access?",
                "Maintain at least two cloud-only emergency accounts with strong phishing-resistant credentials, no dependency on normal federation, narrow exclusions from policies that could lock out all admins, no routine use, and immediate alerting. Store credentials securely, test on a schedule, review logs after every test, and rotate after real use.",
            ),
        ],
    },
    {
        "title": "Azure and Terraform",
        "definition": (
            "Terraform manages Azure through declarative configuration and provider APIs. "
            "Engineers describe desired resources, Terraform compares configuration and state "
            "with Azure, produces a plan, and applies approved changes. Enterprise Terraform "
            "requires more than .tf files: module contracts, remote state, authentication, "
            "testing, release management, drift handling, imports, policy checks, and clear "
            "ownership are essential."
        ),
        "principles": [
            "Separate state by lifecycle and blast radius; do not put an entire enterprise in one state file.",
            "Treat state as sensitive operational data and protect access, network paths, locking, versioning, and recovery.",
            "Pin Terraform, providers, modules, and CI actions; upgrade through reviewed releases.",
            "Use modules as stable product interfaces, not as thin wrappers around every resource.",
            "Use OIDC or managed identity instead of static Azure credentials.",
            "Plan on every change, apply the reviewed commit through an approval boundary, and detect drift.",
        ],
        "building_blocks": [
            ("Configuration", "Root modules compose providers, variables, modules, resources, data sources, checks, and outputs."),
            ("State", "Maps Terraform addresses to remote object identity and stores values required for future plans."),
            ("Providers", "Translate Terraform resource operations into Azure or other service API calls."),
            ("Modules", "Reusable contracts that encapsulate a capability, validation, defaults, and outputs."),
            ("Backends", "Store and lock state; the azurerm backend uses Blob Storage and leases."),
            ("CI/CD", "Formatting, validation, linting, security, tests, plan, approval, apply, and release evidence."),
            ("Import and moved blocks", "Bring existing resources under management and preserve addresses during refactoring."),
            ("Testing", "Static checks, native Terraform tests, contract tests, integration deployments, and policy tests."),
        ],
        "implementation": [
            "Bootstrap state and deployment identity separately with minimal credentials.",
            "Design state boundaries around platform, environment, subscription, service, and team ownership.",
            "Build versioned modules with typed variables, validation, secure defaults, useful outputs, and examples.",
            "Validate formatting and configuration, scan for misconfiguration and secrets, and test representative plans.",
            "Use separate plan and apply identities and protected environments.",
            "Define import, drift, state recovery, provider upgrade, deprecation, and module release procedures.",
        ],
        "pitfalls": [
            "Editing state manually instead of using import, moved, removed, or state commands under controlled recovery.",
            "Using Terraform workspaces as the only isolation for environments with different permissions and lifecycles.",
            "Passing provider-derived unknown values into modules whose data loading requires values at plan time.",
            "Running apply automatically from an untrusted pull request.",
        ],
        "engineer_qa": [
            (
                "Why is Terraform state sensitive?",
                "State contains resource IDs, topology, outputs, and sometimes values returned by providers that may include secrets even when an output is marked sensitive. It also controls what Terraform believes it owns. Encrypt it, use Entra authorization, least privilege, network restrictions, locking, versioning, and recovery; never commit it to Git.",
            ),
            (
                "How do you handle resources that already exist?",
                "Write the desired configuration, verify the exact Azure resource identity, back up state, and use import blocks or `terraform import`. Review the first plan carefully and adjust configuration until it shows no unintended changes. Use moved blocks when changing Terraform addresses so refactoring does not recreate resources.",
            ),
            (
                "What belongs in a reusable Azure module?",
                "A coherent capability with a stable interface, such as a private storage account baseline or subscription profile. Include typed inputs, validation, secure defaults, diagnostics, role and network integration points, meaningful outputs, examples, tests, documentation, semantic versions, and an upgrade policy.",
            ),
        ],
        "design_qa": [
            (
                "Design Terraform state for a landing zone factory.",
                "Use one protected platform state per tenant/environment and one vending or workload state per subscription or lifecycle boundary. Store state in a dedicated subscription with Entra authentication, network restrictions, leases, versioning, and retention. Use deterministic keys, separate plan/apply identities, and do not allow a workload state to control another subscription.",
            ),
            (
                "How would you promote Terraform changes across environments?",
                "Promote immutable module versions and the same reviewed commit, while environment roots supply different non-secret configuration and separate state. Run tests and a plan in development, deploy to a representative canary, then require approval and a new plan for production. Do not copy and manually edit module implementations per environment.",
            ),
            (
                "How should drift be handled?",
                "Run scheduled read-only plans and route findings to the owning team. Determine whether the portal change was authorized, import or update code if it should become desired state, or apply code to restore it. Avoid blind automatic remediation for destructive or availability-sensitive changes and record emergency changes for reconciliation.",
            ),
        ],
    },
    {
        "title": "Platform Architecture",
        "definition": (
            "Platform architecture defines the shared technical capabilities and operating "
            "model that allow product teams to deliver safely and quickly. A cloud platform "
            "is an internal product, not merely a set of centrally owned resources. It has "
            "consumers, paved roads, APIs or request contracts, service levels, support, "
            "security boundaries, lifecycle management, cost, and measurable adoption."
        ),
        "principles": [
            "Design from consumer journeys: request, provision, deploy, observe, support, recover, and retire.",
            "Offer opinionated paved roads with documented escape hatches instead of one inflexible standard.",
            "Separate control plane, shared data-plane services, and workload ownership.",
            "Minimize coupling and shared blast radius; centralize only capabilities that benefit from it.",
            "Publish SLOs, ownership, support paths, version policy, cost, and deprecation timelines.",
            "Use product metrics: lead time, success rate, adoption, compliance, incident load, and consumer satisfaction.",
        ],
        "building_blocks": [
            ("Resource foundation", "Tenant structure, subscriptions, identity, network, security, governance, and state."),
            ("Developer interfaces", "Portal, service catalog, repository template, API, CLI, or versioned request files."),
            ("Delivery platform", "CI/CD, artifact repositories, environments, approvals, policy checks, and deployment identities."),
            ("Runtime platform", "Compute, container, integration, data, secret, and network patterns."),
            ("Observability", "Logs, metrics, traces, SLOs, alerts, dashboards, and incident context."),
            ("Operations", "Ownership, on-call, capacity, backup, patching, recovery, change, and support."),
            ("FinOps", "Unit costs, budgets, allocation, optimization, reservations, and product cost transparency."),
        ],
        "implementation": [
            "Interview product teams and map high-friction delivery and operational journeys.",
            "Define platform capabilities, boundaries, teams, service ownership, and golden paths.",
            "Build a minimum platform slice that delivers one workload securely from request to operations.",
            "Create reusable modules and templates with automated tests and versioned contracts.",
            "Add service catalog discovery, documentation, examples, and support feedback.",
            "Measure adoption and outcomes, then retire low-value customization and reduce bottlenecks.",
        ],
        "pitfalls": [
            "Building a platform without product-team research and then measuring only resource count.",
            "Centralizing every service so the platform team becomes a deployment ticket queue.",
            "Hiding security and cost decisions inside templates with no explanation or override process.",
            "Offering many unmaintained golden paths instead of a small set of reliable products.",
        ],
        "engineer_qa": [
            (
                "What is a platform product?",
                "A platform product is a supported internal service with defined users, outcomes, interfaces, reliability, security, documentation, ownership, cost, and roadmap. For example, subscription vending is a product when teams can submit a validated request and receive a governed subscription with predictable outputs and support—not when an engineer manually runs a script.",
            ),
            (
                "What is the difference between a control plane and a data plane?",
                "The control plane creates and configures resources, policies, identities, and desired state. The data plane carries or stores application data and serves runtime traffic. Access, availability, logging, and network paths often differ, so designs must protect both; Azure Resource Manager access does not automatically grant access to a storage blob or Key Vault secret.",
            ),
            (
                "How do you measure platform success?",
                "Measure workload onboarding lead time, deployment success, change failure, recovery time, platform SLOs, compliance, exception age, unit cost, support demand, adoption, and consumer satisfaction. A platform that is technically standardized but bypassed by teams is not successful.",
            ),
        ],
        "design_qa": [
            (
                "Design an internal developer platform for Azure.",
                "Provide discoverable workload templates and a service catalog that invoke versioned Terraform modules through OIDC pipelines. Integrate subscription vending, approved compute/data patterns, identity, network profiles, secrets, observability, budgets, and ownership. Offer preview plans, protected production environments, outputs usable by application pipelines, SLOs, documentation, and a feedback-driven roadmap.",
            ),
            (
                "What should be centralized versus delegated?",
                "Centralize high-leverage shared foundations such as tenant governance, identity standards, network transit, security evidence, module contracts, and organization-wide observability patterns. Delegate application resources, deployment cadence, service-level tuning, and workload operations within guardrails. Evaluate blast radius, expertise, economies of scale, latency, autonomy, and support capacity.",
            ),
            (
                "How do you avoid the platform team becoming a bottleneck?",
                "Replace tickets with self-service contracts, automate policy and evidence, delegate safe scopes, publish paved roads, and make exceptions visible and time-bound. Design APIs and modules around consumer needs, invest in documentation and diagnostics, and measure queue time and failed requests as platform defects.",
            ),
        ],
    },
    {
        "title": "Azure Services",
        "definition": (
            "Azure services are managed and infrastructure capabilities delivered across "
            "compute, containers, storage, databases, integration, analytics, AI, identity, "
            "networking, management, and security. Enterprise engineering is not about naming "
            "the most services; it is about selecting the least complex service that satisfies "
            "functional, reliability, security, data, operational, regional, and cost requirements."
        ),
        "principles": [
            "Begin with workload requirements and quality attributes, not a preferred product.",
            "Prefer managed services when the reduced operational burden is worth constraints, cost, and platform dependency.",
            "Verify regional availability, zone support, quotas, limits, SLAs, backup, and recovery behavior.",
            "Design identity, networking, DNS, diagnostics, data protection, and cost for every selected service.",
            "Understand control-plane and data-plane access separately.",
            "Use architecture decision records to capture alternatives and trade-offs.",
        ],
        "building_blocks": [
            ("Compute", "Virtual Machines, VM Scale Sets, App Service, Functions, Container Apps, and AKS offer increasing abstraction choices."),
            ("Storage", "Blob, Files, Queues, Tables, managed disks, and Data Lake address object, file, messaging, and disk needs."),
            ("Databases", "Azure SQL, SQL Managed Instance, PostgreSQL, MySQL, Cosmos DB, Redis, and analytics stores have different models and operations."),
            ("Integration", "Service Bus, Event Grid, Event Hubs, Logic Apps, and API Management support messaging, events, workflows, streaming, and APIs."),
            ("Observability", "Azure Monitor, Log Analytics, Application Insights, Managed Prometheus, and Managed Grafana cover telemetry."),
            ("Identity and secrets", "Entra ID, managed identities, Key Vault, App Configuration, and Managed HSM support trust and configuration."),
            ("Management", "Resource Manager, Policy, Resource Graph, Automation, Update Manager, Backup, Site Recovery, and Cost Management."),
            ("Security", "Defender, Sentinel, Firewall, WAF, DDoS Protection, Bastion, and Private Link provide layered controls."),
        ],
        "implementation": [
            "Capture traffic, latency, availability, consistency, data, compliance, recovery, growth, and team skill requirements.",
            "Create a shortlist and compare service limits, dependencies, private networking, identity, observability, and total cost.",
            "Build a representative proof of concept for uncertain performance or integration assumptions.",
            "Define a service baseline module with secure defaults and operational outputs.",
            "Run failure, scaling, backup, restore, and upgrade tests.",
            "Review the choice as Azure capabilities and workload requirements evolve.",
        ],
        "pitfalls": [
            "Choosing AKS for every workload despite the operational cost of a Kubernetes platform.",
            "Using an SLA as if it were the application's end-to-end availability guarantee.",
            "Ignoring service quotas and regional feature differences until deployment.",
            "Selecting a database only from familiarity without consistency, scale, access, and recovery analysis.",
        ],
        "engineer_qa": [
            (
                "How do you choose between App Service, Container Apps, AKS, and VMs?",
                "Use App Service for managed web/API hosting with minimal platform operation, Container Apps for managed containerized apps and event-driven scaling, AKS when Kubernetes APIs and ecosystem control justify cluster operations, and VMs for OS-level control or legacy constraints. Compare networking, scaling, deployment, portability, compliance, team skill, and total operational cost.",
            ),
            (
                "When would you use Service Bus, Event Grid, or Event Hubs?",
                "Service Bus is durable enterprise messaging with queues, topics, transactions, ordering options, and dead-lettering. Event Grid distributes discrete events to subscribers. Event Hubs ingests high-throughput ordered streams for telemetry and analytics. Some architectures use all three for different semantic needs.",
            ),
            (
                "What must an enterprise service baseline include?",
                "Naming and tags, identity, RBAC, private or controlled network access, encryption and key decisions, diagnostics, alerts, backup and recovery, locks where justified, policy compatibility, quotas, cost controls, outputs, tests, and a support owner. A resource deployment without operations is incomplete.",
            ),
        ],
        "design_qa": [
            (
                "Select services for an event-driven order platform.",
                "Use API Management and a suitable web/API compute service for ingress, Service Bus for durable business commands, Event Grid for lightweight state-change notifications, Event Hubs for high-volume telemetry, and a database chosen from transaction and consistency needs. Add managed identities, Key Vault, private access, distributed tracing, dead-letter operations, idempotency, replay strategy, and regional recovery.",
            ),
            (
                "How would you design a service catalog?",
                "Publish supported service patterns with intended use, architecture, security baseline, cost model, limits, SLO, module version, examples, and owner. Let teams request or instantiate patterns through validated interfaces, report adoption and failures, and deprecate versions with migration guidance.",
            ),
            (
                "How do you evaluate a preview Azure feature?",
                "Check support terms, SLA, regional availability, API stability, limits, security and compliance acceptance, data durability, migration path, and provider support. Isolate it behind an abstraction where practical, use noncritical workloads, define an exit strategy, and obtain explicit risk acceptance before production.",
            ),
        ],
    },
    {
        "title": "Hub-and-Spoke Networking",
        "definition": (
            "Hub-and-spoke is an Azure network topology in which shared connectivity and "
            "security services reside in a hub VNet while workload VNets connect as spokes. "
            "VNet peering is nontransitive, so the hub provides routing or inspection between "
            "spokes, hybrid networks, and the internet. The model supports centralized control "
            "and workload isolation but requires careful routing, DNS, scale, and ownership design."
        ),
        "principles": [
            "Keep workload resources in spokes and shared transit, DNS, firewall, and gateways in connectivity subscriptions.",
            "Use explicit user-defined routes when traffic must traverse Azure Firewall or an NVA.",
            "Design the return path and SNAT behavior to prevent asymmetric routing.",
            "Use regional hubs for latency, scale, residency, and failure-domain needs.",
            "Automate peering, route, DNS, and policy attachment as a workload network profile.",
            "Evaluate Virtual WAN when managed transit and connection scale reduce operational burden.",
        ],
        "building_blocks": [
            ("Hub VNet", "Hosts shared transit and approved connectivity services."),
            ("Spoke VNets", "Provide workload isolation, address space, subnets, NSGs, and delegated ownership."),
            ("VNet peering", "Provides low-latency private connectivity but does not itself create transitive routing."),
            ("Azure Firewall or NVA", "Forwards and inspects spoke, hybrid, and egress traffic."),
            ("Gateway transit", "Allows spokes to use a hub VPN or ExpressRoute gateway under peering constraints."),
            ("Route tables", "Steer traffic to firewall, gateways, or other next hops."),
            ("DNS", "Private Resolver and private zones provide centralized resolution and hybrid forwarding."),
        ],
        "implementation": [
            "Create regional address plans for hubs, gateway subnets, firewall subnets, resolver subnets, and spoke growth.",
            "Deploy zone-resilient firewall, gateways, and DNS resolvers where supported.",
            "Create bidirectional peering with forwarded traffic and gateway options set intentionally.",
            "Associate spoke route tables and verify effective routes for internet, hybrid, and other spoke destinations.",
            "Link or centrally manage Private DNS zones and forwarding rules.",
            "Test cross-spoke, hybrid, egress, failover, throughput, and firewall logging before onboarding production.",
        ],
        "pitfalls": [
            "Expecting peering to provide transitive routing without a forwarding appliance or managed hub.",
            "Using one undersized global hub that creates latency, quota, and failure-domain concentration.",
            "Enabling gateway transit and UDRs without understanding propagated routes.",
            "Allowing product teams to create independent private DNS zones for the same namespace.",
        ],
        "engineer_qa": [
            (
                "Why is VNet peering called nontransitive?",
                "If spoke A peers with a hub and spoke B peers with the same hub, A cannot reach B merely because both peerings exist. Traffic must traverse a routing service such as Azure Firewall/NVA, or use a managed topology such as Virtual WAN. Explicit routes and forwarded-traffic settings are required.",
            ),
            (
                "What is gateway transit?",
                "Gateway transit lets a spoke use a VPN or ExpressRoute gateway in a peered hub. The hub peering allows gateway transit and the spoke peering uses remote gateways; a spoke cannot use multiple remote gateways in the same way as independent direct gateways. Route propagation and UDR interaction must be tested.",
            ),
            (
                "Hub-spoke or Virtual WAN?",
                "Traditional hub-spoke offers detailed VNet-level control and fits existing firewall/NVA patterns but requires more routing and peering operations. Virtual WAN provides managed hubs, scalable branch and VNet connectivity, routing intent, and global transit, with its own feature, cost, and customization constraints. Choose from scale, operations, routing, security appliance, and migration requirements.",
            ),
        ],
        "design_qa": [
            (
                "Design a multi-region hub-and-spoke platform.",
                "Deploy a connectivity hub per active region with zonal firewall, DNS resolver, and required gateways. Attach local spokes to reduce latency, connect hubs through global peering or the selected WAN design, define regional and failover routes, centralize policy and logs, and test loss of a hub. Avoid automatically extending every route globally when residency or segmentation forbids it.",
            ),
            (
                "How would you isolate production from nonproduction?",
                "Use separate subscriptions and spokes at minimum, with NSGs and firewall policies. For stronger boundaries, use separate hubs or routing domains and deployment identities, especially when operations, compliance, or hybrid routes differ. Shared services must expose only explicit ports and identities, and logs should prove denied cross-environment traffic.",
            ),
            (
                "How does a new spoke get onboarded safely?",
                "A validated profile allocates nonoverlapping CIDR, creates or enrolls the VNet, configures peering, route tables, NSGs, DNS links, firewall policy references, diagnostics, and ownership. Automated tests verify effective routes, DNS, egress, hybrid access, and forbidden paths before the spoke is marked ready.",
            ),
        ],
    },
    {
        "title": "Private Endpoints",
        "definition": (
            "An Azure private endpoint is a network interface with a private IP address in a "
            "consumer VNet that connects to a specific Azure Private Link service instance. "
            "Clients use the normal service name, DNS resolves it to the private endpoint IP, "
            "and traffic enters the Microsoft service privately. A private endpoint does not "
            "automatically disable public access, configure every DNS client, or secure the "
            "application identity."
        ),
        "principles": [
            "Treat DNS as part of the private endpoint resource lifecycle, not as a later troubleshooting task.",
            "Disable or restrict public network access separately after private-path validation.",
            "Keep authentication and least-privilege data-plane authorization; network privacy is not authorization.",
            "Place endpoints based on consumer access, ownership, scale, and regional failure requirements.",
            "Centralize private DNS zone ownership to prevent duplicate zones and inconsistent records.",
            "Monitor endpoint approval, connection state, DNS records, network reachability, and service logs.",
        ],
        "building_blocks": [
            ("Private endpoint NIC", "Receives a private IP in the consumer VNet and targets one service subresource."),
            ("Private Link service", "Provides private endpoint access to a supported PaaS resource or privately published service."),
            ("Subresources", "Services expose targets such as blob, file, vault, SQL server, or registry with separate DNS needs."),
            ("Private DNS zone", "Maps the service's private-link namespace to endpoint IP addresses."),
            ("Zone links and resolver", "Make records available to Azure VNets and on-premises clients."),
            ("Connection approval", "Resource owners can approve, reject, and monitor endpoint connection requests."),
            ("Public network setting", "Remains a separate service configuration that must be disabled or restricted."),
        ],
        "implementation": [
            "Confirm the service, subresource, region, consumers, authentication, and public access requirements.",
            "Create or reuse the centrally owned Private DNS zone and link the correct resolution VNets.",
            "Create the endpoint in a subnet with sufficient addresses and an explicit ownership model.",
            "Approve the connection where manual approval is required and verify the connection state.",
            "Validate public service name resolution and data-plane access from Azure, on-premises, and recovery locations.",
            "Disable public access, monitor failures, and document endpoint and DNS cleanup ordering.",
        ],
        "pitfalls": [
            "Testing the private IP directly instead of the service hostname, masking TLS and DNS problems.",
            "Creating duplicate zones with the same namespace in separate subscriptions or VNets.",
            "Assuming an NSG alone controls all private endpoint behavior without checking current network policy support and routes.",
            "Running out of subnet addresses because endpoint growth and multiple subresources were not planned.",
        ],
        "engineer_qa": [
            (
                "Private endpoint versus service endpoint?",
                "A private endpoint gives one service instance a private IP in the consumer VNet and supports Private Link DNS. A service endpoint keeps the service's public endpoint and extends VNet identity over the Azure backbone so the service firewall can allow selected subnets. Private endpoints provide stronger instance-level private exposure and broader access patterns, but add DNS and endpoint lifecycle complexity.",
            ),
            (
                "Why does a private endpoint require DNS changes?",
                "Applications normally connect to a public service hostname and validate its TLS certificate. DNS must return the endpoint's private IP while preserving that hostname. Private DNS zones, VNet links, and hybrid forwarding make this split-horizon resolution work consistently.",
            ),
            (
                "Does a private endpoint prevent data exfiltration?",
                "Not by itself. It provides a private path to a specific service instance, but identities, DNS, routes, public egress, and access to unapproved service instances still matter. Use service firewalls or public-access disablement, least privilege, controlled egress, Azure Policy, approved endpoint workflows, and data controls as a combined exfiltration strategy.",
            ),
        ],
        "design_qa": [
            (
                "Design Private Link for centralized services used by many spokes.",
                "Keep service ownership separate from consumer endpoint ownership. Use centrally managed private DNS zones and resolver paths, create endpoints in consumer or shared endpoint VNets based on routing and blast radius, automate records and approvals, allocate address capacity, and expose only required subresources. Avoid a single endpoint when regional availability or consumer isolation requires separate endpoints.",
            ),
            (
                "How would on-premises clients access an Azure private endpoint?",
                "Provide VPN or ExpressRoute reachability to the endpoint VNet, and conditionally forward the service's private-link DNS namespace from on-premises DNS to Azure DNS Private Resolver inbound endpoints. Verify return routes, NSGs and firewall policies, the service hostname, identity authorization, and failure behavior if the Azure region or hybrid link is unavailable.",
            ),
            (
                "How do private endpoints work in a multi-region application?",
                "Create endpoints in or near each consuming regional network and use service-supported regional or global recovery patterns. DNS and application failover must direct clients to a healthy service path; some services need separate regional instances and zones. Test endpoint, DNS, service, and region failures rather than assuming the PaaS replication automatically changes client routing.",
            ),
        ],
    },
]


CROSS_DOMAIN_QA = [
    (
        "Design a secure Azure foundation for a company moving 100 applications from two datacenters.",
        "Begin with discovery: application dependencies, data classification, owners, recovery, IP/DNS, and identity. Deploy an ALZ with platform and workload separation, redundant hybrid connectivity, regional hubs, shared DNS and monitoring, audit-first policy, OIDC Terraform automation, and request-driven subscription vending. Migrate in waves by dependency group, test private endpoints and DNS, measure compliance and performance, and maintain rollback until each workload is operationally accepted.",
    ),
    (
        "A team requests a single large subscription for 30 applications. How do you respond?",
        "Ask about ownership, lifecycle, environments, data boundaries, quotas, network needs, and cost accountability. A subscription is an access, policy, quota, billing, and blast-radius boundary; unrelated applications usually deserve separate subscriptions or at least a deliberate shared-platform model. Offer automated vending so correct boundaries do not create excessive lead time.",
    ),
    (
        "A deny policy breaks a production deployment. What do you do?",
        "Stop repeated deployment attempts, identify the assignment, effect, scope, alias, and exact denied field, and confirm whether the resource is unsafe or the policy has an error. Use the documented emergency process to narrow or temporarily disable enforcement if business impact justifies it, preserve evidence, remediate the workload or policy, validate in a canary, and reconcile all emergency changes through code.",
    ),
    (
        "How do you design for regional failure without doubling every service blindly?",
        "Classify workloads by business impact, RTO, RPO, dependency, and data consistency. Select active-active, active-passive, backup/restore, or redeploy patterns per capability; identify global and regional dependencies such as DNS, identity, Key Vault, container registries, state, and private endpoints. Automate recovery, test it, and measure whether achieved recovery meets the objective.",
    ),
    (
        "How would you review an Azure architecture in an interview or design board?",
        "Start with business goals and quality attributes. Walk through tenant and subscription boundaries, identity, trust zones, data flows, DNS and routing, service selection, availability, security, operations, delivery, cost, and lifecycle. State assumptions, quantify limits, explain trade-offs, identify failure modes, and show how evidence and tests validate the design.",
    ),
    (
        "What should happen when a workload subscription is decommissioned?",
        "Verify owner approval, legal retention, backups, locks, dependencies, private endpoints, DNS, routes, identities, keys, and billing. Remove workload resources through their owning code, preserve required logs and state, revoke access and network paths, move the empty subscription to Decommissioned, observe the retention period, and then cancel through the billing process with auditable approval.",
    ),
    (
        "How do you balance central governance and product-team autonomy?",
        "Central teams define and automate mandatory risk, identity, network, evidence, and cost boundaries. Product teams receive delegated subscriptions and self-service modules within those controls, plus a transparent exception process. Measure both compliance and delivery friction; a control that is constantly bypassed or requires manual tickets needs product and design improvement.",
    ),
    (
        "What is your approach to Azure cost architecture?",
        "Create ownership and allocation metadata at subscription vending, require budgets and alerts, estimate shared platform unit costs, and expose cost data to product teams. Match resilience and service tiers to business objectives, manage commitments from stable usage, detect anomalies, remove expired resources, and include network, logging, security, backup, and data transfer costs—not only compute.",
    ),
    (
        "How do you prove that a cloud design works?",
        "Use multiple evidence types: Terraform validation and tests, policy compliance, effective route and DNS tests, identity access tests, security scans, load and scale tests, backup restoration, regional or link failover, monitoring alerts, incident exercises, and cost measurements. A diagram is a hypothesis until representative failure and operational tests confirm it.",
    ),
    (
        "What are the first questions you ask when requirements are unclear?",
        "Ask who the users and owners are, what data is handled, where it may reside, required availability and recovery, expected scale and latency, integration and network sources, identity types, regulatory obligations, deployment frequency, operational skill, budget, timeline constraints, and exit or decommission expectations. Record assumptions and decisions rather than silently choosing defaults.",
    ),
]


GLOSSARY = [
    ("ALZ", "Azure Landing Zones: Microsoft's architecture and implementation guidance for governed Azure foundations."),
    ("BGP", "Border Gateway Protocol, used to exchange routes between Azure and external networks."),
    ("CIDR", "Classless Inter-Domain Routing notation used to define IP address ranges."),
    ("Control plane", "APIs used to create, configure, and govern resources."),
    ("Data plane", "APIs and network paths used to read, write, or process service data."),
    ("IaC", "Infrastructure as Code: versioned declarative or scripted infrastructure management."),
    ("NVA", "Network virtual appliance, such as a third-party firewall or router."),
    ("OIDC", "OpenID Connect, used for short-lived federated authentication from systems such as GitHub Actions."),
    ("PIM", "Privileged Identity Management for eligible and time-bound privileged access."),
    ("RPO", "Recovery Point Objective: acceptable amount of data loss measured in time."),
    ("RTO", "Recovery Time Objective: acceptable duration to restore service."),
    ("SLA", "Service Level Agreement, a provider commitment under specified conditions."),
    ("SLO", "Service Level Objective, an internal reliability target for a service or platform."),
    ("UDR", "User-defined route, used to override or supplement Azure system routing."),
    ("VNet", "Azure Virtual Network, a private layer-3 network boundary."),
]


def build_document():
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Enterprise Cloud Engineer")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Azure Concepts, Architecture, and Design Q&A")
    document.add_paragraph()
    topics = document.add_paragraph()
    topics.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = topics.add_run(
        "Azure Landing Zones • Governance • Networking • Hybrid Cloud • Security\n"
        "Terraform • Platform Architecture • Azure Services • Hub-and-Spoke • Private Endpoints"
    )
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    purpose = document.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
    purpose.add_run(
        "\nA detailed learning, implementation, architecture review, and interview guide"
    ).italic = True
    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "Perspective: Enterprise Cloud Engineer | Updated: August 2026"
    )
    document.add_page_break()

    document.add_heading("How to Use This Guide", level=1)
    document.add_paragraph(
        "Read the concept explanation first, then use the engineer questions to test "
        "operational understanding and the design questions to practice trade-off discussions. "
        "Strong enterprise answers begin with requirements and assumptions, explain boundaries "
        "and data flows, state alternatives, address failure and operations, and identify evidence."
    )
    add_bullets(
        document,
        [
            "For learning: read one concept and draw its control and data flows from memory.",
            "For implementation: use the principles, building blocks, steps, and pitfalls as a design checklist.",
            "For interviews: answer aloud using Context → Requirements → Design → Trade-offs → Operations → Evidence.",
            "For architecture reviews: adapt the cross-domain scenarios to the organization's real constraints.",
        ],
    )
    document.add_heading("Table of Contents", level=1)
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    document.add_paragraph(
        "If the table is blank, open the file in Microsoft Word, right-click the table, "
        "and select Update Field → Update entire table."
    )
    document.add_page_break()

    document.add_heading("Executive Architecture View", level=1)
    document.add_paragraph(
        "The concepts in this guide form one system. Landing zones establish the tenant and "
        "subscription foundation. Governance defines allowed state and accountability. Networking "
        "and hybrid connectivity move traffic. Security protects identities, control planes, data "
        "planes, and operations. Terraform makes the desired platform repeatable. Platform "
        "architecture packages these capabilities into products. Azure services host workloads, "
        "hub-and-spoke supplies shared transit, and private endpoints provide private PaaS access."
    )
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Layer", "Primary responsibility", "Typical evidence"]
    for index, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], text, bold=True, color=WHITE)
        shade(table.rows[0].cells[index], BLUE)
    rows = [
        ("Business and product", "Ownership, criticality, data, cost, SLO, lifecycle", "Service catalog, RACI, SLO, budget"),
        ("Foundation", "Tenant, management groups, subscriptions, shared platform", "ALZ plan, hierarchy, assignments"),
        ("Governance and security", "Allowed state, access, detection, response", "Compliance, PIM, alerts, exceptions"),
        ("Connectivity", "IP, DNS, routes, ingress, egress, hybrid, private access", "Flow tests, effective routes, DNS tests"),
        ("Workload services", "Compute, data, integration, resilience", "ADRs, service tests, recovery evidence"),
        ("Delivery and operations", "IaC, CI/CD, state, monitoring, support, recovery", "Plans, tests, logs, runbooks, exercises"),
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_data):
            set_cell_text(cells[index], text)
        if len(table.rows) % 2 == 1:
            for cell in cells:
                shade(cell, LIGHT_GRAY)
    document.add_page_break()

    question_number = 1
    for chapter_number, concept in enumerate(CONCEPTS, start=1):
        document.add_heading(f"{chapter_number}. {concept['title']}", level=1)
        document.add_heading("Concept in detail", level=2)
        document.add_paragraph(concept["definition"])

        document.add_heading("Core principles", level=2)
        add_bullets(document, concept["principles"])

        document.add_heading("Main building blocks", level=2)
        building_table = document.add_table(rows=1, cols=2)
        building_table.style = "Table Grid"
        building_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_text(building_table.rows[0].cells[0], "Building block", bold=True, color=WHITE)
        set_cell_text(building_table.rows[0].cells[1], "Purpose", bold=True, color=WHITE)
        shade(building_table.rows[0].cells[0], BLUE)
        shade(building_table.rows[0].cells[1], BLUE)
        for name, description in concept["building_blocks"]:
            cells = building_table.add_row().cells
            set_cell_text(cells[0], name, bold=True)
            set_cell_text(cells[1], description)

        document.add_heading("Enterprise implementation approach", level=2)
        add_numbered(document, concept["implementation"])

        document.add_heading("Common design mistakes", level=2)
        add_bullets(document, concept["pitfalls"])

        document.add_heading("Enterprise Cloud Engineer Q&A", level=2)
        for question, answer in concept["engineer_qa"]:
            add_qa(document, question, answer, question_number)
            question_number += 1

        document.add_heading("Design Q&A", level=2)
        for question, answer in concept["design_qa"]:
            add_qa(document, question, answer, question_number)
            question_number += 1

        if chapter_number != len(CONCEPTS):
            document.add_page_break()

    document.add_page_break()
    document.add_heading("11. Cross-Domain Enterprise Design Q&A", level=1)
    document.add_paragraph(
        "These questions combine multiple concepts. A strong answer should identify assumptions, "
        "show control and data flows, discuss failure modes and operations, and explain why the "
        "selected design is preferable to reasonable alternatives."
    )
    for question, answer in CROSS_DOMAIN_QA:
        add_qa(document, question, answer, question_number)
        question_number += 1

    document.add_page_break()
    document.add_heading("12. Enterprise Architecture Review Checklist", level=1)
    checklists = {
        "Business and ownership": [
            "Named service owner, technical owner, security owner, cost owner, and support path",
            "Data classification, residency, regulatory, retention, and privacy requirements",
            "Availability SLO, RTO, RPO, scale, performance, and budget",
            "Provisioning, change, incident, recovery, and decommission journeys",
        ],
        "Foundation and governance": [
            "Management group, subscription, resource group, and environment boundaries",
            "Policy scope, parameters, effects, remediation identities, exemptions, and evidence",
            "RBAC least privilege, PIM, workload identity, emergency access, and access reviews",
            "Tags, budgets, naming, locks, quotas, and ownership metadata",
        ],
        "Networking and hybrid": [
            "Nonoverlapping IP plan, DNS resolution, route propagation, and return paths",
            "Ingress, east-west, egress, management, and hybrid flows",
            "Firewall, NSG, WAF, DDoS, private endpoint, and public access decisions",
            "Capacity, zones, regions, provider diversity, failover, and observability",
        ],
        "Security and data": [
            "Threat model, trust boundaries, authentication, authorization, and secret elimination",
            "Encryption, key ownership, rotation, recovery, backup, and immutability",
            "Defender coverage, security logs, detections, incidents, and response automation",
            "Supply-chain, code, dependency, container, IaC, and artifact controls",
        ],
        "Delivery and operations": [
            "Terraform state boundary, provider/module pinning, imports, and drift",
            "OIDC identity, plan/apply separation, approvals, tests, and release evidence",
            "Logs, metrics, traces, alerts, SLO dashboards, and on-call ownership",
            "Failure exercises, capacity tests, restoration, rollback, and deprecation",
        ],
    }
    for heading, items in checklists.items():
        document.add_heading(heading, level=2)
        for item in items:
            document.add_paragraph(f"☐ {item}")

    document.add_page_break()
    document.add_heading("13. Interview Answer Framework", level=1)
    document.add_paragraph(
        "Use this sequence for design questions. It demonstrates engineering judgment instead "
        "of jumping directly to a product name."
    )
    framework = [
        ("1. Clarify", "Users, owners, data, scale, latency, compliance, availability, RTO/RPO, budget, and current state."),
        ("2. State assumptions", "Make missing facts explicit and explain what would change the design."),
        ("3. Define boundaries", "Tenant, subscription, management group, identity, network, region, and team ownership."),
        ("4. Draw flows", "Control plane, user/data traffic, DNS, routes, identity tokens, logs, state, and recovery."),
        ("5. Select capabilities", "Choose services from requirements and compare at least one credible alternative."),
        ("6. Secure and govern", "Least privilege, policy, secrets, private/public access, data protection, detection, and exceptions."),
        ("7. Design failure", "Zones, regions, dependencies, capacity, RTO/RPO, backup, failover, and degraded modes."),
        ("8. Operate", "IaC, deployment, monitoring, ownership, SLOs, support, cost, patching, and decommission."),
        ("9. Validate", "Plans, policy evidence, access tests, DNS/routes, load, restoration, failover, and security exercises."),
        ("10. Explain trade-offs", "Cost, complexity, autonomy, consistency, lock-in, performance, and migration path."),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.rows[0].cells[0], "Step", bold=True, color=WHITE)
    set_cell_text(table.rows[0].cells[1], "What to cover", bold=True, color=WHITE)
    shade(table.rows[0].cells[0], BLUE)
    shade(table.rows[0].cells[1], BLUE)
    for step, detail in framework:
        cells = table.add_row().cells
        set_cell_text(cells[0], step, bold=True)
        set_cell_text(cells[1], detail)

    document.add_page_break()
    document.add_heading("14. Glossary", level=1)
    glossary_table = document.add_table(rows=1, cols=2)
    glossary_table.style = "Table Grid"
    set_cell_text(glossary_table.rows[0].cells[0], "Term", bold=True, color=WHITE)
    set_cell_text(glossary_table.rows[0].cells[1], "Meaning", bold=True, color=WHITE)
    shade(glossary_table.rows[0].cells[0], BLUE)
    shade(glossary_table.rows[0].cells[1], BLUE)
    for term, meaning in GLOSSARY:
        cells = glossary_table.add_row().cells
        set_cell_text(cells[0], term, bold=True)
        set_cell_text(cells[1], meaning)

    document.add_heading("Closing perspective", level=2)
    document.add_paragraph(
        "Enterprise Cloud Engineers connect strategy to implementation. They understand Azure "
        "services, but their greater value is designing safe boundaries, explaining trade-offs, "
        "automating repeatable outcomes, making data and control flows visible, and proving that "
        "the platform can be operated and recovered. Use this guide as a starting point, then "
        "validate every design against current Microsoft documentation, service limits, regional "
        "availability, organizational risk, and real workload tests."
    )

    core_properties = document.core_properties
    core_properties.title = "Enterprise Cloud Engineer — Azure Concepts, Architecture, and Design Q&A"
    core_properties.subject = "Detailed Azure concepts and enterprise cloud engineering design questions and answers"
    core_properties.author = "Azure Landing Zone Factory"
    core_properties.keywords = (
        "Azure, Landing Zones, Governance, Networking, Hybrid Cloud, Security, "
        "Terraform, Platform Architecture, Hub-and-Spoke, Private Endpoints"
    )

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
